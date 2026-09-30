#!/usr/bin/env python3
"""Validate a generated .pptx deck: integrity, bounds, overflow, transitions.

    python validate_deck.py deck.pptx [--expect 12] [--json]

Errors   = the file is broken or unreadable as intended (bad zip, shape off the
           slide, overlapping text/tables/charts, a table that grows off the
           slide, malformed or out-of-order transition XML).
Warnings = it will open fine but probably looks wrong (text likely overflows its
           box, a table will auto-grow, tiny type, empty slide, repeated title,
           no transitions at all).

Exit code is 1 when there is at least one error, so this is usable as a gate in
a script or a test. Tables are measured, not skipped: PowerPoint grows a row to
fit its text, so a table can be taller than its declared box. Every estimate here
is a heuristic shared with the builder (see textmetrics.py) - it catches gross
layout failure, it does not typeset.

Known blind spot: a run with no explicit font size inherits one from the slide
master, which python-pptx cannot resolve without reading the theme. Such runs are
not measured, so decks built on a foreign template with `--base` are checked for
bounds and transitions only.
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.enum.text import PP_ALIGN

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from textmetrics import (  # noqa: E402  - sys.path above makes this importable
    CELL_INSET,
    CELL_PAD,
    est_lines,
    row_heights,
    text_width,
)

EMU_IN = 914400.0
MIN_FONT_PT = 10.0

NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
VALID_TRANSITIONS = {"fade", "dissolve", "cut", "push", "wipe", "cover", "split"}
# CT_Slide child order: anything after transition must be one of these
AFTER_TRANSITION = {"timing", "extLst"}


def paragraph_size(p):
    for run in p.runs:
        if run.font.size is not None:
            return run.font.size.pt
    if p.font.size is not None:
        return p.font.size.pt
    return None


def is_chrome(shape):
    """Footer, source and slide-number shapes are author-chosen 9pt furniture."""
    return shape.name.startswith("chrome:")


def is_bleed(shape):
    """Shapes named `bleed:...` are meant to run off the slide edge on purpose."""
    return shape.name.startswith("bleed:")


def text_metrics(shape):
    """(min_font_pt, required_height_in) for a shape's visible text."""
    if not shape.has_text_frame or shape.width is None:
        return None, 0.0
    tf = shape.text_frame
    width = shape.width / EMU_IN - (tf.margin_left + tf.margin_right) / EMU_IN
    if width <= 0.1:
        return None, 0.0
    min_font, needed = None, 0.0
    for p in tf.paragraphs:
        text = "".join(r.text for r in p.runs) or (p.text or "")
        if not text.strip():
            continue
        size = paragraph_size(p)
        if size is None:
            continue
        min_font = size if min_font is None else min(min_font, size)
        needed += est_lines(text, width, size) * size * 1.22 / 72.0
        if p.space_after is not None:
            needed += p.space_after.pt / 72.0
        if p.space_before is not None:
            needed += p.space_before.pt / 72.0
    return min_font, needed


def table_metrics(shape):
    """(declared_h, required_h, min_font_pt, grid) in inches / points.

    PowerPoint treats a declared row height as a *minimum* and grows the row until
    its text fits, so the table a reader sees can be taller than the box the
    author placed - which is how a table ends up sitting on the footer. Re-deriving
    each row's required height is what lets us notice that before it ships.
    """
    table = shape.table
    widths = [c.width / EMU_IN for c in table.columns]
    declared = sum(r.height for r in table.rows) / EMU_IN
    probe = table.cell(0, 0)
    inset = (probe.margin_left + probe.margin_right) / EMU_IN
    pad = (probe.margin_top + probe.margin_bottom) / EMU_IN
    grid, sizes, min_font = [], [], None
    for row in table.rows:
        texts, size = [], 0.0
        for cell in row.cells:
            texts.append(cell.text or "")
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    if run.font.size is not None:
                        size = max(size, run.font.size.pt)
        grid.append(texts)
        sizes.append(size or 12.0)
        min_font = size if min_font is None else min(min_font, size)
    required = sum(row_heights(grid, widths, sizes, pad=pad or CELL_PAD,
                               inset=inset or CELL_INSET))
    return declared, required, min_font, grid


def ink_box(shape):
    """The rectangle a shape actually paints into, in inches.

    Declared box width is not ink width: a left-aligned text box is routinely
    much wider than the words inside it, so comparing raw boxes invents collisions
    no reader can see. Estimating the ink, honouring right/centre alignment, and
    using a table's *grown* height is what makes a collision report worth acting
    on. Returns (x0, y0, x1, y1, label) or None.
    """
    if shape.width is None or shape.height is None:
        return None
    left = (shape.left or 0) / EMU_IN
    top = (shape.top or 0) / EMU_IN
    w, h = shape.width / EMU_IN, shape.height / EMU_IN
    if shape.has_table:
        _, required, _, _ = table_metrics(shape)
        return left, top, left + w, top + max(required, h), "table %r" % shape.name
    if shape.has_chart:
        return left, top, left + w, top + h, "chart %r" % shape.name
    if not shape.has_text_frame:
        return None
    tf = shape.text_frame
    box = w - (tf.margin_left + tf.margin_right) / EMU_IN
    if box <= 0.1:
        return None
    _, needed_h = text_metrics(shape)
    if needed_h <= 0:
        return None
    widest, alignment = 0.0, None
    for p in tf.paragraphs:
        text = "".join(r.text for r in p.runs) or (p.text or "")
        if not text.strip():
            continue
        size = paragraph_size(p)
        if size is None:
            continue
        width = text_width(text, size)
        if width > widest:
            widest, alignment = width, p.alignment
    if widest <= 0:
        return None
    x0 = left + tf.margin_left / EMU_IN
    y0 = top + tf.margin_top / EMU_IN
    widest = min(widest, box)
    if alignment == PP_ALIGN.RIGHT:
        x0 += box - widest
    elif alignment == PP_ALIGN.CENTER:
        x0 += (box - widest) / 2.0
    label = (tf.text or "").strip().replace("\n", "|")[:32] or shape.name
    return x0, y0, x0 + widest, y0 + min(h, max(needed_h, 0.06)), repr(label)


def check_collisions(slide):
    """Shapes whose painted rectangles overlap. These are errors, not warnings.

    Overlapping text is not a style opinion - it is unreadable - and a table or
    chart frame under a text block hides that text. Tolerances are the width of
    the estimate, not permission to collide.
    """
    boxes = [b for b in (ink_box(sh) for sh in slide.shapes) if b]
    problems = []
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            x0a, y0a, x1a, y1a, la = boxes[i]
            x0b, y0b, x1b, y1b, lb = boxes[j]
            ox = min(x1a, x1b) - max(x0a, x0b)
            oy = min(y1a, y1b) - max(y0a, y0b)
            opaque = la.startswith(("table ", "chart ")) or lb.startswith(("table ", "chart "))
            if ox > (0.02 if opaque else 0.10) and oy > (0.04 if opaque else 0.08):
                problems.append("shapes collide by %.2fx%.2fin: %s over %s"
                                % (ox, oy, la, lb))
    return problems


def check_degenerate(shape):
    """A shape with no extent at all renders as nothing.

    This is the signature of freeform coordinates being rounded away - python-pptx
    derives a freeform's extent from its integer path coordinates, so passing
    inches straight in collapses anything under an inch.
    """
    w = (shape.width or 0) / EMU_IN
    h = (shape.height or 0) / EMU_IN
    if w > 0.002 or h > 0.002:
        return None
    return "shape %r has zero extent and will not be visible" % (shape.name,)


def check_text_frame(shape, slide_w, slide_h):
    """Return (warnings, min_font_seen, text_chars) for one shape."""
    warnings, min_font, chars = [], None, 0
    if not shape.has_text_frame:
        return warnings, min_font, chars
    tf = shape.text_frame
    if shape.width is None or shape.height is None:
        return warnings, min_font, chars
    w_in = shape.width / EMU_IN - (tf.margin_left + tf.margin_right) / EMU_IN
    h_in = shape.height / EMU_IN - (tf.margin_top + tf.margin_bottom) / EMU_IN
    if w_in <= 0.1 or h_in <= 0.1:
        return warnings, min_font, chars
    needed = 0.0
    for p in tf.paragraphs:
        text = "".join(r.text for r in p.runs) or (p.text or "")
        chars += len(text)
        if not text.strip():
            continue
        size = paragraph_size(p)
        if size is None:
            continue
        min_font = size if min_font is None else min(min_font, size)
        if size < MIN_FONT_PT and not is_chrome(shape):
            warnings.append("%.1fpt text is below the %.0fpt legibility floor: %r"
                            % (size, MIN_FONT_PT, text[:40]))
        lead = size * 1.22
        lines = est_lines(text, w_in, size)
        needed += lines * lead / 72.0
        if p.space_after is not None:
            needed += p.space_after.pt / 72.0
        if p.space_before is not None:
            needed += p.space_before.pt / 72.0
    if needed > h_in * 1.06 and needed > 0:
        warnings.append("text may overflow: needs ~%.2fin in a %.2fin box (%r...)"
                        % (needed, h_in, (tf.text or "")[:40].replace("\n", " ")))
    return warnings, min_font, chars


def _qn(tag):
    prefix, local = tag.split(":")
    uri = {"p": NS_P, "p14": NS_P14, "mc": NS_MC, "r": NS_R}.get(prefix)
    return "{%s}%s" % (uri, local)


def slide_parts_in_order(zf):
    """Slide parts in display order, resolved through the presentation rels."""
    names = set(zf.namelist())
    pres = etree.fromstring(zf.read("ppt/presentation.xml"))
    rels = etree.fromstring(zf.read("ppt/_rels/presentation.xml.rels"))
    target = {rel.get("Id"): rel.get("Target") for rel in rels}
    parts = []
    for sld_id in pres.iter(_qn("p:sldId")):
        tgt = target.get(sld_id.get(_qn("r:id")))
        if not tgt:
            continue
        tgt = tgt[1:] if tgt.startswith("/") else "ppt/" + tgt.lstrip("./")
        tgt = tgt.replace("\\", "/")
        if tgt in names:
            parts.append(tgt)
    return parts


def check_transitions(zf):
    """Inspect embedded <p:transition> elements for placement and validity."""
    result = {"count": 0, "slides": [], "errors": [], "warnings": []}
    try:
        parts = slide_parts_in_order(zf)
    except (KeyError, etree.XMLSyntaxError) as exc:
        result["errors"].append("cannot resolve slide order: %s" % exc)
        return result
    for i, part in enumerate(parts, start=1):
        try:
            root = etree.fromstring(zf.read(part))
        except etree.XMLSyntaxError as exc:
            result["errors"].append("%s is not valid XML: %s" % (part, exc))
            continue
        trans = root.find(_qn("p:transition"))
        if trans is None:
            continue
        result["count"] += 1
        entry = {"slide": i, "type": None, "duration_ms": None}
        # the effect element must be the only child
        kids = [etree.QName(c).localname for c in trans]
        if len(kids) != 1 or kids[0] not in VALID_TRANSITIONS:
            result["errors"].append(
                "%s: transition effect must be exactly one of %s, found %s"
                % (part, sorted(VALID_TRANSITIONS), kids))
        else:
            entry["type"] = kids[0]
        dur = trans.get(_qn("p14:dur"))
        if dur is not None:
            if not dur.isdigit():
                result["errors"].append("%s: p14:dur is not an integer (%r)" % (part, dur))
            else:
                entry["duration_ms"] = int(dur)
            ignorable = (root.get(_qn("mc:Ignorable")) or "").split()
            if "p14" not in ignorable:
                result["errors"].append(
                    "%s: uses p14:dur but mc:Ignorable does not list p14; "
                    "pre-2010 readers would reject the slide" % part)
        idx = list(root).index(trans)
        trailing = [etree.QName(c).localname for c in list(root)[idx + 1:]]
        stray = [t for t in trailing if t not in AFTER_TRANSITION]
        if stray:
            result["errors"].append(
                "%s: <p:transition> is out of schema order, followed by %s"
                % (part, stray))
        result["slides"].append(entry)
    if not result["count"]:
        result["warnings"].append("no slide transitions are embedded in this deck")
    return result


def validate(path, expect=None):
    report = {"file": path, "ok": False, "slide_count": 0, "slides": [],
              "errors": [], "warnings": []}

    try:
        with zipfile.ZipFile(path) as zf:
            bad = zf.testzip()
        if bad:
            report["errors"].append("zip contains a corrupt entry: %s" % bad)
            return report
    except (zipfile.BadZipFile, OSError) as exc:
        report["errors"].append("not a readable .pptx (zip) file: %s" % exc)
        return report

    try:
        prs = Presentation(path)
    except Exception as exc:                                  # noqa: BLE001
        report["errors"].append("python-pptx cannot open it: %s" % exc)
        return report

    slide_w = prs.slide_width or 0
    slide_h = prs.slide_height or 0
    report["slide_width_in"] = round(slide_w / EMU_IN, 3)
    report["slide_height_in"] = round(slide_h / EMU_IN, 3)
    if slide_w / EMU_IN < 12.0:
        report["warnings"].append("slide is not 16:9 (%.2fin wide)"
                                  % (slide_w / EMU_IN))

    with zipfile.ZipFile(path) as zf:
        report["transitions"] = check_transitions(zf)

    slides = list(prs.slides)
    report["slide_count"] = len(slides)
    if expect is not None and len(slides) != expect:
        report["errors"].append("expected %d slides, found %d" % (expect, len(slides)))

    seen_titles = {}
    for i, slide in enumerate(slides, start=1):
        entry = {"index": i, "title": "", "shapes": len(slide.shapes),
                 "errors": [], "warnings": []}
        text_chars = 0
        min_font = None
        for shape in slide.shapes:
            left = shape.left if shape.left is not None else 0
            top = shape.top if shape.top is not None else 0
            w = shape.width or 0
            h = shape.height or 0
            tol = int(0.02 * EMU_IN)
            if (not is_bleed(shape)) and (left + w > slide_w + tol
                                          or top + h > slide_h + tol
                                          or left < -tol or top < -tol):
                entry["errors"].append(
                    "shape %r extends past the slide edge" % (shape.name,))
            if shape.has_table:
                declared, required, cell_font, _ = table_metrics(shape)
                for row in shape.table.rows:
                    for cell in row.cells:
                        text_chars += len(cell.text or "")
                if cell_font is not None:
                    min_font = cell_font if min_font is None else min(min_font, cell_font)
                    if cell_font < MIN_FONT_PT and not is_chrome(shape):
                        entry["warnings"].append(
                            "%.1fpt table text is below the %.0fpt legibility floor"
                            % (cell_font, MIN_FONT_PT))
                if required > declared + 0.02:
                    entry["warnings"].append(
                        "table rows declare %.2fin but the cells need ~%.2fin; "
                        "PowerPoint will grow it past where it was placed"
                        % (declared, required))
                grown = (shape.top or 0) / EMU_IN + required
                if grown > slide_h / EMU_IN:
                    entry["errors"].append(
                        "table grows %.2fin past the bottom of the slide"
                        % (grown - slide_h / EMU_IN))
                continue
            if shape.has_chart:
                continue
            warns, mf, chars = check_text_frame(shape, slide_w, slide_h)
            text_chars += chars
            entry["warnings"].extend(warns)
            if mf is not None:
                min_font = mf if min_font is None else min(min_font, mf)
        # the title is the largest type on the slide, not the first text found
        # (otherwise every slide would look like it repeated its running kicker)
        best_size, best_text = -1.0, ""
        for shape in slide.shapes:
            if is_chrome(shape) or not shape.has_text_frame:
                continue
            for p in shape.text_frame.paragraphs:
                text = "".join(r.text for r in p.runs) or (p.text or "")
                text = text.strip()
                if not text:
                    continue
                size = paragraph_size(p) or 0.0
                if size > best_size:
                    best_size, best_text = size, text
        entry["title"] = best_text[:60]
        entry["min_font_pt"] = min_font
        entry["errors"].extend(check_collisions(slide))
        for shape in slide.shapes:
            degenerate = check_degenerate(shape)
            if degenerate:
                entry["warnings"].append(degenerate)
        if text_chars == 0:
            entry["warnings"].append("slide has no text at all")
        key = entry["title"].lower()
        if key:
            seen_titles.setdefault(key, []).append(i)
        report["slides"].append(entry)

    for title, where in seen_titles.items():
        if len(where) > 1:
            report["warnings"].append("title %r repeats on slides %s"
                                      % (title[:50], where))
    report["errors"].extend(e for s in report["slides"] for e in s["errors"])
    report["warnings"].extend(w for s in report["slides"] for w in s["warnings"])
    report["errors"].extend(report["transitions"]["errors"])
    report["warnings"].extend(report["transitions"]["warnings"])
    report["ok"] = not report["errors"]
    return report


def print_report(r):
    print("Deck: %s" % r["file"])
    if "slide_width_in" in r:
        print("Size: %.2f x %.2f in" % (r["slide_width_in"], r["slide_height_in"]))
    print("Slides: %d" % r["slide_count"])
    print("")
    for s in r["slides"]:
        status = "ERROR" if s["errors"] else ("WARN" if s["warnings"] else "ok")
        print("  %2d  %-38s shapes=%-3d %s"
              % (s["index"], s["title"][:38] or "(no text)", s["shapes"], status))
        for msg in s["errors"]:
            print("        ERROR: %s" % msg)
        for msg in s["warnings"]:
            print("        WARN:  %s" % msg)
    print("")
    for msg in r["errors"]:
        if not any(msg in s["errors"] for s in r["slides"]):
            print("ERROR: %s" % msg)
    for msg in r["warnings"]:
        if not any(msg in s["warnings"] for s in r["slides"]):
            print("WARN:  %s" % msg)
    tr = r.get("transitions")
    if tr:
        kinds = {}
        for s in tr["slides"]:
            kinds[s["type"]] = kinds.get(s["type"], 0) + 1
        detail = ", ".join("%s x%d" % (k, v) for k, v in sorted(kinds.items()))
        print("Transitions: %d embedded%s" % (tr["count"], " (%s)" % detail if detail else ""))
    print("")
    print("Summary: %d slides, %d errors, %d warnings"
          % (r["slide_count"], len(r["errors"]), len(r["warnings"])))
    print("Result: %s" % ("PASS (editable, opens cleanly)" if r["ok"] else "FAIL"))


def main(argv=None):
    ap = argparse.ArgumentParser(description="Validate a .pptx deck.")
    ap.add_argument("deck")
    ap.add_argument("--expect", type=int, help="expected slide count")
    ap.add_argument("--json", action="store_true", help="machine-readable output")
    args = ap.parse_args(argv)
    report = validate(args.deck, expect=args.expect)
    if args.json:
        print(json.dumps(report, indent=2))
    else:
        print_report(report)
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    sys.exit(main())
