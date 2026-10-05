#!/usr/bin/env python3
"""Regression checks for the layout rules that let broken slides ship.

    python tests/test_layout_rules.py

Plain asserts, no test framework: every check below fails loudly if the rule it
guards regresses. Two of them are the bugs this suite was written for - a table
sized for its shortest cell (so PowerPoint grew it over the footer) and a stack
note box allowed to run into the marker column.
"""
import json
import os
import sys
import tempfile
from pathlib import Path

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.join(os.path.dirname(HERE), "scripts")
sys.path.insert(0, SCRIPTS)

import validate_deck                                     # noqa: E402
from build_deck import build, fit_size                             # noqa: E402
from pptx import Presentation                            # noqa: E402
from pptx.util import Pt                                 # noqa: E402
from textmetrics import column_widths, row_heights       # noqa: E402

EMU_IN = 914400.0


def _deck(slides):
    return {"title": "Test Deck", "slides": slides}


def _write(tmp, spec, name="t.pptx"):
    path = os.path.join(tmp, name)
    build(spec, path)
    return path


def _shapes(path, index):
    return list(Presentation(path).slides[index].shapes)


def test_columns_follow_content_demand_with_a_floor():
    """Space goes to the column that holds the most text, and none starve."""
    grid = [["Blocker", "Why it stalled", "What broke it"],
            ["Growing a usable crystal", "GaN films cracked on every substrate",
             "AlN buffer layer"]]
    widths = column_widths(grid, 11.633)
    demand = [max(len(r[c]) for r in grid) for c in range(3)]
    assert abs(sum(widths) - 11.633) < 0.01, widths
    assert all(w >= 1.5 - 1e-9 for w in widths), widths
    assert widths.index(max(widths)) == demand.index(max(demand)), (widths, demand)


def test_rows_get_the_height_of_their_worst_cell():
    widths = [2.0, 3.0]
    grid = [["a", "short"], ["b", "a cell whose text wraps onto several lines of its own"]]
    heights = row_heights(grid, widths, 12)
    assert heights[1] > heights[0], heights


def test_table_does_not_need_growing(tmp_path_placeholder=None):
    """Every declared row height must already cover its cell text."""
    with tempfile.TemporaryDirectory() as tmp:
        spec = _deck([{
            "layout": "table", "title": "Three blockers",
            "headers": ["The blocker", "Why it stalled the field", "What broke it open"],
            "rows": [
                ["Growing a usable crystal", "GaN films cracked and came out unusable",
                 "AlN buffer layer - Amano & Akasaki, 1986"],
                ["Making GaN p-type",
                 "Mg-doped GaN stayed insulating; without a p-side there is no junction",
                 "Electron-beam irradiation, 1989; thermal annealing, 1992"],
                ["Getting bright light out", "Early GaN emitters were orders of magnitude too dim",
                 "InGaN double heterostructure - Nakamura, 1993"],
            ]}])
        path = _write(tmp, spec)
        table = next(s.table for s in Presentation(path).slides[0].shapes if s.has_table)
        declared = sum(r.height for r in table.rows) / EMU_IN
        _, required, _, _ = validate_deck.table_metrics(
            next(s for s in Presentation(path).slides[0].shapes if s.has_table))
        assert required <= declared + 0.02, (required, declared)


def test_table_too_dense_to_read_is_refused():
    """A table that cannot fit at 12pt must raise, not shrink below the floor."""
    spec = _deck([{
        "layout": "table", "title": "Too much",
        "headers": ["A", "B", "C"],
        "rows": [["x" * 240, "y" * 240, "z" * 240] for _ in range(9)]}])
    try:
        build(spec, os.path.join(tempfile.gettempdir(), "never-written.pptx"))
    except ValueError as exc:
        assert "12pt" in str(exc), exc
    else:
        raise AssertionError("an unreadably dense table was accepted")


def test_stack_note_cannot_reach_the_marker_column():
    with tempfile.TemporaryDirectory() as tmp:
        spec = _deck([{
            "layout": "stack", "title": "The InGaN junction",
            "items": [
                {"label": "p-type GaN", "note": "Magnesium doping supplies the holes",
                 "marker": "h+"},
                {"label": "InGaN active layer",
                 "note": "Electrons and holes recombine here, emitting a blue photon",
                 "marker": "450 nm"},
            ]}])
        boxes = [s for s in _shapes(_write(tmp, spec), 0)
                 if s.has_text_frame and s.text_frame.text.strip()]
        notes = sorted((b.left for b in boxes if "recombine" in b.text_frame.text))
        markers = sorted((b.left for b in boxes if b.text_frame.text.strip() == "450 nm"))
        note_right = notes[0] + next(b.width for b in boxes
                                     if "recombine" in b.text_frame.text)
        assert note_right <= markers[0], (note_right / EMU_IN, markers[0] / EMU_IN)


def test_validator_reports_collisions_as_errors():
    """The gate must fail a deck whose text genuinely collides."""
    with tempfile.TemporaryDirectory() as tmp:
        spec = _deck([{
            "layout": "bullets", "title": "Collide",
            "bullets": ["A" * 120]}])
        path = _write(tmp, spec)
        prs = Presentation(path)
        slide = prs.slides[0]
        victim = max((s for s in slide.shapes if s.has_text_frame and s.text_frame.text),
                     key=lambda s: len(s.text_frame.text))
        ghost = slide.shapes.add_textbox(victim.left, victim.top, victim.width,
                                         victim.height)
        ghost.text_frame.text = "overlapping text of a similar length to the original"
        # An explicit size matters: text that inherits its size from a master is
        # not measurable, which is itself a limitation of the estimates.
        ghost.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
        prs.save(path)
        report = validate_deck.validate(path)
        assert any("collide" in e for e in report["errors"]), report["errors"]


def test_clean_deck_passes_the_gate():
    with tempfile.TemporaryDirectory() as tmp:
        spec = _deck([
            {"layout": "title", "title": "A Deck", "subtitle": "That is fine"},
            {"layout": "bullets", "title": "A modest slide", "bullets": ["One.", "Two."]},
        ])
        report = validate_deck.validate(_write(tmp, spec))
        assert report["ok"], report["errors"] + report["warnings"]


def test_text_that_cannot_fit_is_refused():
    try:
        fit_size(["word " * 1000], 2, 1)
    except ValueError as exc:
        assert "split" in str(exc)
    else:
        raise AssertionError("overflowing text was silently accepted")


def test_image_alt_text_bounds_and_collisions():
    from PIL import Image
    with tempfile.TemporaryDirectory() as tmp:
        image = os.path.join(tmp, "figure.png")
        Image.new("RGB", (800, 300), "white").save(image)
        path = _write(tmp, _deck([{
            "layout": "image", "title": "Evidence", "image": image,
            "alt": "A labelled timeline", "caption": "Conceptual diagram"}]))
        report = validate_deck.validate(path)
        assert report["ok"], report["errors"]
        assert not any("alt text" in w for w in report["warnings"])
        prs = Presentation(path)
        picture = next(s for s in prs.slides[0].shapes
                       if s.shape_type == validate_deck.MSO_SHAPE_TYPE.PICTURE)
        picture._element.nvPicPr.cNvPr.set("descr", "")
        ghost = prs.slides[0].shapes.add_textbox(picture.left, picture.top,
                                               picture.width, picture.height)
        ghost.text_frame.text = "This text is hidden by the picture"
        ghost.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
        prs.save(path)
        report = validate_deck.validate(path)
        assert any("alt text" in w for w in report["warnings"])
        assert any("collide" in e for e in report["errors"])


def test_user_reference_composition_uses_existing_layouts():
    """The supplied problem/solution + journey treatment remains opt-in."""
    with tempfile.TemporaryDirectory() as tmp:
        path = _write(tmp, _deck([
            {"layout": "two_column", "title": "Problem and response",
             "left": {"heading": "Problem", "bullets": ["Access is slow", "Quality is uncertain"]},
             "right": {"heading": "Proposed solution", "bullets": ["Assess quickly", "Track each batch"]}},
            {"layout": "sequence", "title": "One connected journey",
             "steps": ["Select", "Assess", "Plan", "Order", "Value"],
             "highlight": 4, "note": "Example workflow; adapt stages to the real service."},
        ]))
        prs = Presentation(path)
        assert len(prs.slides) == 2
        assert any("Problem" in sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
        assert any("Select" in sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame)
        assert validate_deck.validate(path)["ok"]


def test_bar_axis_and_background_are_explicit():
    with tempfile.TemporaryDirectory() as tmp:
        path = _write(tmp, {"title": "Test", "theme": {"bg": "F8F5EF"}, "slides": [{
            "layout": "chart", "title": "Measured", "unit": "years",
            "categories": ["A", "B"], "series": [{"name": "Age", "values": [5, 9]}]}]})
        slide = Presentation(path).slides[0]
        assert str(slide.background.fill.fore_color.rgb) == "F8F5EF"
        chart = next(s.chart for s in slide.shapes if s.has_chart)
        assert chart.value_axis.minimum_scale == 0
        assert chart.value_axis.has_title


def test_generated_and_template_decks_are_transition_free():
    from lxml import etree
    with tempfile.TemporaryDirectory() as tmp:
        path = _write(tmp, _deck([{"layout": "title", "title": "Open in Office"}]))
        report = validate_deck.validate(path)
        assert report["ok"] and report["transitions"]["count"] == 0
        assert not any("no slide transitions" in w for w in report["warnings"])
        base = Presentation(path)
        transition = etree.SubElement(base.slides[0]._element,
                                     "{http://schemas.openxmlformats.org/presentationml/2006/main}transition")
        etree.SubElement(transition, "{http://schemas.openxmlformats.org/presentationml/2006/main}fade")
        base.save(path)
        original = Path(path).read_bytes()
        report = validate_deck.validate(path)
        assert not report["ok"] and report["transitions"]["count"] == 1
        assert any("not allowed" in e for e in report["errors"])
        output = os.path.join(tmp, "clean.pptx")
        build(_deck([]), output, base=path)
        assert validate_deck.validate(output)["transitions"]["count"] == 0
        assert Path(path).read_bytes() == original


def main():
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    failed = 0
    for test in tests:
        try:
            test()
        except AssertionError as exc:
            failed += 1
            print("FAIL %s: %s" % (test.__name__, exc))
        else:
            print("ok   %s" % test.__name__)
    print("\n%d/%d passed" % (len(tests) - failed, len(tests)))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
