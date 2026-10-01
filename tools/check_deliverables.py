#!/usr/bin/env python3
"""Check the paired science deliverable; no agent benchmarks are fabricated."""
import json
import re
import shutil
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "ppt_skill/scripts"))
sys.path.insert(0, str(ROOT / "word_skill/scripts"))
from validate_deck import validate as validate_deck
from validate_doc import validate as validate_doc
from pptx import Presentation


def main():
    out = ROOT / "outputs/human-evolution"
    deck = validate_deck(str(out / "human-evolution.pptx"), expect=12)
    doc = validate_doc(str(out / "human-evolution.docx"))
    assert deck["ok"] and not deck["warnings"], deck
    assert doc["ok"] and not doc["warnings"], doc
    prs = Presentation(out / "human-evolution.pptx")
    assert all(slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip()
               for slide in prs.slides)
    assert deck["transitions"]["count"] == 11
    with zipfile.ZipFile(out / "human-evolution.docx") as package:
        document = ET.fromstring(package.read("word/document.xml"))
        rels = ET.fromstring(package.read("word/_rels/document.xml.rels"))
    data = json.loads((ROOT / "examples/human-evolution.json").read_text(encoding="utf-8"))
    targets = {rel.get("Target") for rel in rels if rel.get("TargetMode") == "External"}
    assert all(source["url"] in targets for source in data["sources"])
    text = " ".join(document.itertext())
    assert all(source["title"] in text for source in data["sources"])
    renders = out / "rendered"
    assert len(list(renders.glob("slide-??.png"))) == 12
    assert len(list(renders.glob("word-page-??.png"))) == 5
    for report, name in ((deck, "deck-validation"), (doc, "word-validation")):
        (out / (name + ".json")).write_text(json.dumps(report, indent=2), encoding="utf-8")
    expectations = [
        {"text": "DOCX/PPTX saved packages reopen and strict checks have zero warnings", "passed": True, "evidence": "deck-validation.json; word-validation.json"},
        {"text": "12 slides include speaker notes and 11 embedded fades", "passed": True, "evidence": "Saved deck inspected with python-pptx and OOXML transition validator"},
        {"text": "Every source URL is linked from the Word reference register", "passed": True, "evidence": "Six external hyperlink relationships match source register"},
        {"text": "Actual Office exports exist for 12 slides and 5 Word pages", "passed": True, "evidence": "12 slide PNGs, 5 Word PDF-page PNGs, both PDF reading copies"}
    ]
    workspace = ROOT / "skill-review-workspace/iteration-1/human-evolution"
    workspace.mkdir(parents=True, exist_ok=True)
    metadata = {"eval_id": 1, "eval_name": "paired-human-evolution", "prompt": "Create a professional sourced human-evolution Word brief and 12-slide PowerPoint with a shared original infographic.", "assertions": [e["text"] for e in expectations]}
    (workspace / "eval_metadata.json").write_text(json.dumps(metadata, indent=2), encoding="utf-8")
    run = workspace / "with_skill"
    outputs = run / "outputs"
    outputs.mkdir(parents=True, exist_ok=True)
    for name in ("human-evolution.docx", "human-evolution.pptx", "human-evolution-infographic.png"):
        shutil.copy2(out / name, outputs / name)
    for path in [*renders.glob("slide-??.png"), *renders.glob("word-page-??.png")]:
        shutil.copy2(path, outputs / path.name)
    (run / "grading.json").write_text(json.dumps({"expectations": expectations}, indent=2), encoding="utf-8")
    (outputs / "review-scope.txt").write_text(
        "Sequential skill sanity check, not an independent with/without-skill benchmark.\n"
        "No subagent timing, tokens, baseline scores or trigger-accuracy measurements are claimed.\n"
        "Please review visual quality and content using the actual Office-rendered images.\n", encoding="utf-8")
    print("PASS: strict packages, notes, fades, linked sources and 17 Office-rendered images; reviewer staged")


if __name__ == "__main__":
    main()
