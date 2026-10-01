#!/usr/bin/env python3
"""Structural DOCX checks; use Word/LibreOffice rendering for pagination QA."""
import argparse
import json
import zipfile
from docx import Document
from docx.oxml.ns import qn


def validate(path):
    report = {"file": str(path), "errors": [], "warnings": [], "ok": False}
    try:
        with zipfile.ZipFile(path) as package:
            bad = package.testzip()
            if bad:
                raise ValueError("corrupt entry: " + bad)
        doc = Document(path)
    except Exception as exc:
        report["errors"].append("cannot reopen DOCX: " + str(exc))
        return report
    if not doc.core_properties.title:
        report["warnings"].append("document title metadata is missing")
    if not any(p.style.name == "Heading 1" for p in doc.paragraphs):
        report["errors"].append("no semantic Heading 1 paragraphs")
    if not any(p.text.strip() for p in doc.paragraphs):
        report["errors"].append("document is empty")
    for image in doc.inline_shapes:
        if not (image._inline.docPr.get("descr") or "").strip():
            report["errors"].append("image is missing alt text")
        section = doc.sections[0]
        if image.width > section.page_width - section.left_margin - section.right_margin:
            report["errors"].append("image exceeds the text region")
    for table in doc.tables:
        if table.rows[0]._tr.find(".//" + qn("w:tblHeader")) is None:
            report["warnings"].append("table header will not repeat across pages")
    for p in list(doc.paragraphs) + [p for t in doc.tables for row in t.rows
                                   for c in row.cells for p in c.paragraphs]:
        for run in p.runs:
            if run.font.size is not None and run.font.size.pt < 9:
                report["warnings"].append("explicit text below 9pt")
    report["headings"] = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    report["images"] = len(doc.inline_shapes)
    report["tables"] = len(doc.tables)
    report["ok"] = not report["errors"]
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("document")
    parser.add_argument("--json", action="store_true")
    parser.add_argument("--strict", action="store_true")
    args = parser.parse_args()
    report = validate(args.document)
    print(json.dumps(report, indent=2) if args.json else
          "%s: %d errors, %d warnings (structural; render separately)" %
          ("PASS" if report["ok"] else "FAIL", len(report["errors"]), len(report["warnings"])))
    return 0 if report["ok"] and not (args.strict and report["warnings"]) else 1


if __name__ == "__main__":
    raise SystemExit(main())
