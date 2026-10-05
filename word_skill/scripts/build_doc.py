#!/usr/bin/env python3
"""Build an editable .docx from JSON; image paths are relative to the spec."""
import argparse
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


THEME = {"primary": "163D38", "accent": "176B60", "text": "243530",
         "muted": "53615C", "panel": "E8F0EB", "font": "Calibri"}


def element(tag, **attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(qn(key), str(value))
    return node


def inline(paragraph, text):
    for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", str(text))):
        if chunk:
            paragraph.add_run(chunk).bold = i % 2 == 1


def link(paragraph, label, url):
    if not url.startswith(("https://", "http://")):
        raise ValueError("reference URL must use http or https")
    rel = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = element("w:hyperlink", **{"r:id": rel})
    run = element("w:r")
    props = element("w:rPr")
    props.append(element("w:rStyle", **{"w:val": "Hyperlink"}))
    run.append(props)
    text = element("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def check_spec(spec):
    if not isinstance(spec, dict) or not str(spec.get("title", "")).strip():
        raise ValueError("document title is required")
    if not spec.get("sections") or not isinstance(spec["sections"], list):
        raise ValueError("sections must be a non-empty list")
    for section in spec["sections"]:
        if not section.get("heading") or not section.get("blocks"):
            raise ValueError("each section needs heading and blocks")
        for block in section["blocks"]:
            kind = block.get("type", "paragraph")
            if kind not in {"paragraph", "bullets", "callout", "image", "table", "references"}:
                raise ValueError("unknown block type: %s" % kind)
            if kind == "image" and not block.get("alt", "").strip():
                raise ValueError("images need descriptive alt text")
            if kind == "table":
                headers, rows = block.get("headers", []), block.get("rows", [])
                if not headers or not rows or any(len(row) != len(headers) for row in rows):
                    raise ValueError("table must have headers and rectangular rows")
                if len(headers) > 5:
                    raise ValueError("split tables with more than five columns")


def configure(doc, spec):
    theme = {**THEME, **spec.get("theme", {})}
    for value in (theme[k] for k in ("primary", "accent", "text", "muted", "panel")):
        if not re.fullmatch(r"[0-9A-Fa-f]{6}", value):
            raise ValueError("theme colours must be six-digit hex")
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.72)
    section.left_margin = section.right_margin = Inches(0.78)
    section.header_distance = section.footer_distance = Inches(0.32)
    for name, size, colour in (("Normal", 11, "text"), ("Title", 32, "primary"),
                               ("Subtitle", 14, "muted"), ("Heading 1", 23, "primary"),
                               ("Heading 2", 15, "accent"), ("Caption", 10, "muted")):
        style = doc.styles[name]
        style.font.name, style.font.size = theme["font"], Pt(size)
        style.font.color.rgb = RGBColor.from_string(theme[colour])
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.widow_control = True
        if name.startswith("Heading"):
            style.paragraph_format.keep_with_next = True
    doc.styles["Normal"].paragraph_format.line_spacing = 1.15
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.text = spec.get("running_title", spec["title"])
    header.style = doc.styles["Caption"]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.style = doc.styles["Caption"]
    footer.text = spec.get("footer", "Research brief") + "  |  Page "
    footer._p.append(element("w:fldSimple", **{"w:instr": "PAGE"}))
    doc.core_properties.title = spec["title"]
    doc.core_properties.subject = spec.get("subtitle", "")
    props = doc.core_properties
    props.author = spec.get("author", "")
    props.last_modified_by = ""
    props.keywords = ""
    props.comments = ""
    props.category = ""
    props.revision = 1
    props.created = props.modified = datetime.now(timezone.utc)
    return theme


def build(spec, output, asset_root=None):
    check_spec(spec)
    doc = Document()
    theme = configure(doc, spec)
    doc.add_paragraph(spec["title"], "Title")
    if spec.get("subtitle"):
        doc.add_paragraph(spec["subtitle"], "Subtitle")
    if spec.get("meta"):
        doc.add_paragraph(spec["meta"], "Caption")
    root = Path(asset_root or ".")
    for section in spec["sections"]:
        if section.get("page_break"):
            doc.add_page_break()
        doc.add_heading(section["heading"], level=1)
        for block in section["blocks"]:
            kind = block.get("type", "paragraph")
            if kind == "paragraph":
                inline(doc.add_paragraph(), block["text"])
            elif kind == "bullets":
                for text in block["items"]:
                    inline(doc.add_paragraph(style="List Bullet"), text)
            elif kind == "callout":
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(12)
                p._p.get_or_add_pPr().append(element("w:shd", **{"w:fill": theme["panel"]}))
                inline(p, block["text"])
            elif kind == "image":
                path = Path(block["path"])
                if not path.is_absolute():
                    path = root / path
                width = block.get("width", 6.7)
                if not 0 < width <= 6.71:
                    raise ValueError("image width must fit the 6.71-inch text region")
                picture = doc.add_picture(str(path), width=Inches(width))
                picture._inline.docPr.set("descr", block["alt"])
                doc.paragraphs[-1].paragraph_format.keep_with_next = bool(block.get("caption"))
                if block.get("caption"):
                    doc.add_paragraph(block["caption"], "Caption")
            elif kind == "table":
                headers, rows = block["headers"], block["rows"]
                table = doc.add_table(rows=1, cols=len(headers))
                table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = False
                for column in table.columns:
                    column.width = Inches(6.7 / len(headers))
                table.rows[0]._tr.get_or_add_trPr().append(element("w:tblHeader"))
                for cell, text in zip(table.rows[0].cells, headers):
                    cell.text = str(text)
                    cell._tc.get_or_add_tcPr().append(element("w:shd", **{"w:fill": theme["primary"]}))
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                for i, row in enumerate(rows):
                    for cell, text in zip(table.add_row().cells, row):
                        cell.text = str(text)
                        cell._tc.get_or_add_tcPr().append(element("w:shd", **{
                            "w:fill": theme["panel"] if i % 2 == 0 else "FFFFFF"}))
                for row in table.rows:
                    row._tr.get_or_add_trPr().append(element("w:cantSplit"))
                    for cell in row.cells:
                        cell.width = Inches(6.7 / len(headers))
                        for p in cell.paragraphs:
                            p.paragraph_format.space_after = Pt(5)
            elif kind == "references":
                for item in block["items"]:
                    p = doc.add_paragraph()
                    p.add_run("[%s] " % item["id"]).bold = True
                    link(p, item["title"], item["url"])
                    if item.get("note"):
                        p.add_run("\n" + item["note"])
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return str(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("spec")
    parser.add_argument("-o", "--out", required=True)
    args = parser.parse_args()
    path = Path(args.spec)
    build(json.loads(path.read_text(encoding="utf-8")), args.out, path.resolve().parent)
    print("wrote " + args.out)


if __name__ == "__main__":
    main()
